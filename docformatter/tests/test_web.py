import base64
import io
import json
import os
import re

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.settings import INGRESS_PROXY_IP, Settings
from app.web import creer_app

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
BASE = "/api/hassio_ingress/abc"


def _docx(*paragraphes: str) -> bytes:
    doc = Document()
    for p in paragraphes:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def settings(tmp_path):
    s = Settings(config_dir=tmp_path / "config", share_dir=tmp_path / "share", data_dir=tmp_path / "data")
    s.config_dir.mkdir()
    s.dictionnaires_dir.mkdir()
    return s


@pytest.fixture
def app(settings):
    return creer_app(settings)


def client(app, utilisateur="secretaire", ip=INGRESS_PROXY_IP):
    c = TestClient(app, client=(ip, 50000), follow_redirects=False)
    c.headers.update({"X-Remote-User-Name": utilisateur, "X-Ingress-Path": BASE})
    return c


def texte_docx(c, page) -> list[str]:
    """Paragraphes du premier document téléchargeable de la page de résultat."""
    lien = page.text.split('href="telecharger/')[1].split('"')[0]
    return [p.text for p in Document(io.BytesIO(c.get(f"/telecharger/{lien}").content)).paragraphs]


def formater(c, *paragraphes, nom="cr.docx"):
    r = c.post("/formater", files={"fichiers": (nom, _docx(*paragraphes), DOCX)})
    assert r.status_code == 303 and r.headers["location"].startswith(f"{BASE}/lot/")
    cle = r.headers["location"].rsplit("/", 1)[1]
    return cle, c.get(f"/lot/{cle}")


# ---- Accès ----------------------------------------------------------------------------

def test_refuse_hors_ingress(app):
    assert client(app, ip="192.168.1.20").get("/").status_code == 403


def test_base_ingress(app):
    assert f'<base href="{BASE}/">' in client(app).get("/").text


def test_toutes_les_pages_accessibles(app):
    c = client(app)
    for url in ("/", "/vocabulaire", "/mise-en-forme", "/historique", "/expert"):
        r = c.get(url)
        assert r.status_code == 200, url
        assert 'href="expert"' in r.text
    assert c.get("/expert/api/remplacements").status_code == 200


# ---- Formatage ------------------------------------------------------------------------

def test_formater_et_telecharger(app):
    c = client(app)
    _, page = formater(c, "ATCD:HTA")
    assert page.status_code == 200 and "Antécédents" in page.text
    lien = page.text.split('href="telecharger/')[1].split('"')[0]
    fichier = c.get(f"/telecharger/{lien}")
    assert fichier.status_code == 200 and fichier.content[:2] == b"PK"
    assert "filename*=UTF-8''cr%20-%20format%C3%A9.docx" in fichier.headers["content-disposition"]
    # Ni le document ni le lot ne sont accessibles à un autre utilisateur
    autre = client(app, utilisateur="autre")
    assert autre.get(f"/telecharger/{lien}").status_code == 404


def test_lot_prive(app):
    cle, _ = formater(client(app))
    assert client(app, utilisateur="autre").get(f"/lot/{cle}").status_code == 404


def test_fichier_invalide(app):
    c = client(app)
    r = c.post("/formater", files={"fichiers": ("cr.docx", b"pas un docx", DOCX)})
    assert "pas pu être lu" in c.get(r.headers["location"].removeprefix(BASE)).text
    r = c.post("/formater", files={"fichiers": ("cr.pdf", b"%PDF", "application/pdf")})
    assert "seuls les fichiers Word" in c.get(r.headers["location"].removeprefix(BASE)).text


def test_decisions_depuis_le_document(app, monkeypatch):
    """La secrétaire valide un mot et en corrige un autre ; le document est reformaté."""
    from app.pipeline.model import Inconnu
    from app.pipeline.spelling import Correcteur

    def faux_verifier(self, textes):
        return [[Inconnu(t.find(m), t.find(m) + len(m), m, ["dyspnée"])
                 for m in ("dispnée", "coroscan") if m in t and m.lower() not in self.dictionnaire]
                for t in textes]

    monkeypatch.setattr(Correcteur, "disponible", property(lambda self: True))
    monkeypatch.setattr(Correcteur, "verifier", faux_verifier)

    c = client(app)
    cle, page = formater(c, "Patient avec une dispnée, coroscan normal")
    assert "Mots à vérifier" in page.text and 'value="dispnée"' in page.text

    mots = page.text.split('name="mot-')
    index = {m.split('value="')[1].split('"')[0]: m.split('"')[0] for m in mots[1:]}
    r = c.post(f"/lot/{cle}/decisions", data={
        f"mot-{index['dispnée']}": "dispnée", f"choix-{index['dispnée']}": "remplacer",
        f"par-{index['dispnée']}": "dyspnée",
        f"mot-{index['coroscan']}": "coroscan", f"choix-{index['coroscan']}": "correct",
        f"par-{index['coroscan']}": "",
    })
    assert r.status_code == 303
    page = c.get(f"/lot/{cle}")
    assert "C&#39;est noté" in page.text or "C'est noté" in page.text
    assert "Mots à vérifier" not in page.text  # plus rien à vérifier après reformatage
    assert "dyspnée" in page.text

    vocab = c.get("/api/vocabulaire").json()
    assert {"texte": "dispnée", "par": "dyspnée"} in vocab["remplacements"]
    assert "coroscan" in vocab["mots"]
    assert "Depuis un document" in c.get("/historique").text


def test_decisions_vides(app):
    c = client(app)
    cle, _ = formater(c, "Patient stable")
    c.post(f"/lot/{cle}/decisions", data={"mot-0": "x", "choix-0": ""})
    assert "rien n&#39;a changé" in c.get(f"/lot/{cle}").text


# ---- Vocabulaire ----------------------------------------------------------------------

def test_vocabulaire_et_annulation(app):
    c = client(app)
    r = c.post("/api/vocabulaire/remplacement", json={"texte": "ttt", "par": "traitement"})
    assert r.status_code == 200
    assert {"texte": "ttt", "par": "traitement"} in r.json()["vocabulaire"]["remplacements"]
    annulation = r.json()["annulation"]

    _, page = formater(c, "ttt par IPP")
    assert "<ins>Traitement</ins> par IPP" in page.text

    assert c.post("/api/historique/restaurer", json={"id": annulation, "annulation": True}).status_code == 200
    assert {"texte": "ttt", "par": "traitement"} not in c.get("/api/vocabulaire").json()["remplacements"]
    assert "Annulation : Remplacement ajouté" in c.get("/historique").text


def test_mots_connus(app):
    c = client(app)
    assert c.post("/api/vocabulaire/mot", json={"mot": "Xarelto"}).status_code == 200
    r = c.post("/api/vocabulaire/mot", json={"mot": "xarelto"})
    assert r.status_code == 422 and "déjà" in r.json()["erreur"]
    r = c.post("/api/vocabulaire/mot", json={"mot": "deux mots"})
    assert r.status_code == 422 and "Un seul mot" in r.json()["erreur"]
    r = c.request("DELETE", "/api/vocabulaire/mot", json={"mot": "Xarelto"})
    assert "Xarelto" not in r.json()["vocabulaire"]["mots"]


def test_remplacement_identique_refuse(app):
    r = client(app).post("/api/vocabulaire/remplacement", json={"texte": "HTA", "par": "hta"})
    assert r.status_code == 422 and "identique" in r.json()["erreur"]


def test_rubriques(app):
    c = client(app)
    assert c.post("/api/vocabulaire/rubrique", json={"titre": "Évolution"}).status_code == 200
    assert c.post("/api/vocabulaire/variante", json={"titre": "Évolution", "variante": "evol"}).status_code == 200
    _, page = formater(c, "EVOL : favorable")
    assert "Évolution\u00a0: favorable" in texte_docx(c, page)
    r = c.request("DELETE", "/api/vocabulaire/variante", json={"titre": "Évolution", "variante": "evol"})
    rubrique = next(x for x in r.json()["vocabulaire"]["rubriques"] if x["titre"] == "Évolution")
    assert rubrique["variantes"] == []
    r = c.post("/api/vocabulaire/variante", json={"titre": "Inexistante", "variante": "x"})
    assert r.status_code == 422


# ---- Mise en forme --------------------------------------------------------------------

def test_mise_en_forme(app):
    c = client(app)
    page = c.get("/mise-en-forme").text
    assert "Espacer les unités" in page and "75mg" in page
    r = c.post("/api/mise-en-forme/regle", json={"index": 4, "nom": "Espacer les unités", "actif": False})
    assert r.status_code == 200
    _, resultat = formater(c, "Kardégic 75mg")
    assert "75mg" in resultat.text
    # Index et nom incohérents (règles modifiées entre-temps) → refus
    r = c.post("/api/mise-en-forme/regle", json={"index": 0, "nom": "Espacer les unités", "actif": True})
    assert r.status_code == 422
    assert c.post("/api/mise-en-forme/option", json={"cle": "point_final", "actif": True}).status_code == 200
    assert c.post("/api/mise-en-forme/option", json={"cle": "inconnue", "actif": True}).status_code == 422


# ---- Mode expert ----------------------------------------------------------------------

def test_expert_enregistrer_et_conflit(app):
    c = client(app, utilisateur="autre")
    donnees = c.get("/expert/api/remplacements").json()
    donnees["remplacements"].append({"nom": "Test", "motif": "abc", "remplacement": "abd"})
    r = c.put("/expert/api/remplacements", json=donnees)
    assert r.status_code == 200 and r.json()["remplacements"][-1]["nom"] == "Test"
    # La secrétaire modifie le vocabulaire : l'ancienne version de l'expert est périmée
    client(app).post("/api/vocabulaire/mot", json={"mot": "apixaban2"})
    r = c.put("/expert/api/remplacements", json={"version": r.json()["version"], "remplacements": []})
    assert r.status_code == 409


def test_expert_regex_invalide(app):
    c = client(app, utilisateur="autre")
    donnees = c.get("/expert/api/remplacements").json()
    donnees["remplacements"].append({"nom": "cassée", "motif": "(abc"})
    r = c.put("/expert/api/remplacements", json=donnees)
    assert r.status_code == 422 and "expression régulière invalide" in r.json()["erreur"]


def test_expert_tester(app):
    c = client(app, utilisateur="autre")
    r = c.post("/expert/api/tester-regle", json={"motif": r"(\d)mg", "remplacement": r"\1{nbsp}mg", "exemple": "5mg"})
    assert r.json()["resultat"] == "5\u00a0mg"
    r = c.post("/expert/api/tester-regle", json={"motif": "a", "remplacement": r"\2", "exemple": "a"})
    assert "remplacement invalide" in r.json()["erreur"]
    r = c.post("/expert/api/tester", json={"remplacements": [{"nom": "x", "motif": "stable", "remplacement": "STABLE"}],
                                            "texte": "patient stable"})
    assert "<ins>STABLE</ins>" in r.json()["html"]


def test_import_export(app):
    c = client(app, utilisateur="autre")
    assert "remplacements:" in c.get("/expert/export").text
    r = c.post("/expert/import", files={"fichier": ("r.yaml", "corrections:\n  abc: abd\n", "application/x-yaml")})
    assert r.status_code == 200
    assert c.get("/api/vocabulaire").json()["remplacements"] == [{"texte": "abc", "par": "abd"}]
    r = c.post("/expert/import", files={"fichier": ("r.yaml", "remplacements: [{nom: x, motif: '('}]", "text/yaml")})
    assert r.status_code == 422




def test_alerte_fichier_abime(settings):
    settings.regles_path.write_text("options: [pas, valide", encoding="utf-8")
    app = creer_app(settings)
    assert "Vous pouvez continuer à travailler normalement" in client(app).get("/").text


# ---- Configuration et dossier surveillé -----------------------------------------------

def test_options_addon(tmp_path, monkeypatch):
    from app.settings import charger_settings

    options = tmp_path / "options.json"
    options.write_text(json.dumps({"dossier_surveille": True}))
    monkeypatch.setenv("DOCFORMATTER_OPTIONS", str(options))
    monkeypatch.setenv("DOCFORMATTER_CONFIG", str(tmp_path / "config"))
    monkeypatch.setenv("DOCFORMATTER_DATA", str(tmp_path / "data"))
    s = charger_settings()
    assert s.dossier_surveille


def test_dossier_surveille(tmp_path):
    from app.watcher import Surveillant

    s = Surveillant(tmp_path / "share", lambda data: b"ok")
    fichier = s.entree / "cr.docx"
    fichier.write_bytes(b"x")
    os.utime(fichier, (0, 0))
    s.passe()
    assert (s.sortie / "cr.docx").read_bytes() == b"ok"
    assert (s.traites / "cr.docx").exists()


# ---- Retouche manuelle ----------------------------------------------------------------

def seg(texte, **mise_en_forme):
    return {"texte": texte, **mise_en_forme}


def sans_surlignage(html: str) -> str:
    """Aperçu sans le surlignage des mots inconnus (qui dépend de la présence de Hunspell)."""
    return re.sub(r"<mark[^>]*>|</mark>", "", html)


def editeur(c, cle, n=0):
    return c.get(f"/api/lot/{cle}/document/{n}").json()


def textes(blocs):
    return ["".join(s["texte"] for s in b["segments"]) for b in blocs if b["type"] == "paragraphe"]


def test_retouche_page_et_enregistrement(app):
    c = client(app)
    cle, page = formater(c, "ATCD : HTA", "patient stable")
    assert f'href="lot/{cle}/document/0"' in page.text
    assert c.get(f"/lot/{cle}/document/0").status_code == 200

    donnees = editeur(c, cle)
    blocs = donnees["blocs"]
    assert textes(blocs) == ["Antécédents\u00a0: HTA", "Patient stable"]
    blocs[1]["segments"] = [seg("ttt  à revoir")]
    blocs.append({"type": "paragraphe", "id": None, "origine": blocs[1]["id"], "segments": [seg("Ajout")]})
    r = c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    assert r.status_code == 200
    # Pas de règle automatique sur la saisie (« ttt » et l'espace double restent)
    assert textes(r.json()["blocs"]) == ["Antécédents\u00a0: HTA", "ttt  à revoir", "Ajout"]

    fichier = c.get(f"/telecharger/{r.json()['telechargement']}")
    paragraphes = [p.text for p in Document(io.BytesIO(fichier.content)).paragraphs]
    assert paragraphes == ["Antécédents\u00a0: HTA", "ttt  à revoir", "Ajout"]
    resultat = sans_surlignage(c.get(f"/lot/{cle}").text)
    assert "retouché à la main" in resultat and "ttt  à revoir" in resultat


def test_retouche_gras_italique_souligne(app):
    c = client(app)
    cle, _ = formater(c, "texte")
    blocs = editeur(c, cle)["blocs"]
    blocs[0]["segments"] = [seg("Pression "), seg("très élevée", gras=True), seg(" à "),
                            seg("surveiller", italique=True, souligne=True), seg(".")]
    r = c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    segments = r.json()["blocs"][0]["segments"]
    assert [(s["texte"], s["gras"], s["italique"], s["souligne"]) for s in segments] == [
        ("Pression ", False, False, False), ("très élevée", True, False, False), (" à ", False, False, False),
        ("surveiller", False, True, True), (".", False, False, False),
    ]
    doc = Document(io.BytesIO(c.get(f"/telecharger/{r.json()['telechargement']}").content))
    runs = [(run.text, bool(run.bold), bool(run.italic), bool(run.underline)) for run in doc.paragraphs[0].runs]
    assert runs == [("Pression ", False, False, False), ("très élevée", True, False, False),
                    (" à ", False, False, False), ("surveiller", False, True, True), (".", False, False, False)]
    apercu = sans_surlignage(c.get(f"/lot/{cle}").text)
    assert "<strong>très élevée</strong>" in apercu and "<em><u>surveiller</u></em>" in apercu


def test_retouche_conservee_apres_ajout_de_vocabulaire(app):
    c = client(app)
    cle, _ = formater(c, "texte d'origine")
    blocs = editeur(c, cle)["blocs"]
    blocs[0]["segments"] = [seg("ttt "), seg("retouché", gras=True)]
    c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    # Une nouvelle règle ne réécrit pas la retouche…
    c.post("/api/vocabulaire/remplacement", json={"texte": "ttt", "par": "traitement"})
    assert "ttt <strong>retouché</strong>" in sans_surlignage(c.get(f"/lot/{cle}").text)
    # … sauf un remplacement choisi explicitement depuis le résultat (mise en forme conservée)
    c.post(f"/lot/{cle}/decisions", data={"mot-0": "retouché", "choix-0": "remplacer", "par-0": "revu"})
    assert "ttt <strong>revu</strong>" in sans_surlignage(c.get(f"/lot/{cle}").text)


def test_abandon_retouche(app):
    c = client(app)
    cle, _ = formater(c, "patient stable")
    blocs = editeur(c, cle)["blocs"]
    blocs[0]["segments"] = [seg("autre chose")]
    c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    r = c.delete(f"/api/lot/{cle}/document/0")
    assert textes(r.json()["blocs"]) == ["Patient stable"]
    assert "retouché à la main" not in c.get(f"/lot/{cle}").text


def test_retouche_refusee(app):
    c = client(app)
    cle, _ = formater(c, "patient stable")
    blocs = editeur(c, cle)["blocs"]
    blocs[0]["segments"] = [seg("  ")]
    r = c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    assert r.status_code == 422 and "vide" in r.json()["erreur"]
    assert c.put(f"/api/lot/{cle}/document/0", json={"blocs": [{"type": "script"}]}).status_code == 422
    r = c.put(f"/api/lot/{cle}/document/0", json={"blocs": [{"type": "paragraphe", "id": "p42", "segments": []}]})
    assert r.status_code == 422 and "rechargez" in r.json()["erreur"]
    assert c.get(f"/lot/{cle}/document/5").status_code == 404
    assert client(app, utilisateur="autre").get(f"/lot/{cle}/document/0").status_code == 404


# ---- Collage depuis Word --------------------------------------------------------------

def test_coller_affiche_le_recapitulatif_puis_l_editeur(app):
    c = client(app)
    titre = {"alignement": "centre", "police": "Arial", "taille": 12, "espace_apres": 6}
    texte = {"retrait_gauche": 36, "retrait_premiere_ligne": -18, "interligne": 1.07, "police": "Arial", "taille": 10}
    r = c.post("/api/coller", json={"blocs": [
        {"segments": [seg("COMPTE-RENDU", gras=True, souligne=True)], "mise_en_page": titre},
        {"segments": [], "mise_en_page": {"police": "Arial", "taille": 10}},
        {"segments": [seg("ATCD", gras=True), seg(" : HTA , diabéte\tsuivi")], "mise_en_page": texte},
        {"type": "tableau", "bordures": True, "lignes": [[[{"segments": [seg("Créatinine")]}],
                                                          [{"segments": [seg("96µmol/L")]}]]]},
    ]})
    assert r.status_code == 200
    adresse = r.json()["adresse"]
    assert re.fullmatch(r"lot/[\w-]+", adresse)
    cle = adresse.split("/")[1]
    # Récapitulatif : corrections détaillées, « Copier pour Word » avec les blocs du document
    recap = c.get(f"/{adresse}").text
    assert "Compte rendu collé" in recap and "Détail des corrections" in recap
    assert 'class="principal copier"' in recap and 'id="blocs-0"' in recap
    blocs_recap = json.loads(recap.split('id="blocs-0" type="application/json">')[1].split("</script>")[0])
    assert blocs_recap == editeur(c, cle)["blocs"]
    page = c.get(f"/{adresse}/document/0")
    assert page.status_code == 200 and "Relire le compte rendu" in page.text and 'id="copier"' in page.text

    blocs = editeur(c, cle)["blocs"]
    # Règles appliquées (rubrique, espaces, remplacement), ligne vide gardée, tabulation conservée
    assert textes(blocs) == ["COMPTE-RENDU", "", "Antécédents\u00a0: HTA, diabète\tsuivi"]
    assert blocs[0]["segments"] == [{"texte": "COMPTE-RENDU", "gras": True, "italique": False, "souligne": True}]
    assert blocs[0]["mise_en_page"] == {**titre, "espace_avant": 0}
    assert blocs[1]["mise_en_page"] == {"police": "Arial", "taille": 10, "espace_avant": 0, "espace_apres": 0}
    assert blocs[2]["mise_en_page"] == {**texte, "espace_avant": 0, "espace_apres": 0}
    tableau = blocs[3]
    assert tableau["mise_en_page"] == {"bordures": True}
    assert [["".join(s["texte"] for s in p["segments"]) for p in cellule] for cellule in tableau["lignes"][0]] == [
        ["Créatinine"], ["96\u00a0µmol/L"]]


def test_coller_nouveau_paragraphe_reprend_la_mise_en_page(app):
    c = client(app)
    adresse = c.post("/api/coller", json={"blocs": [
        {"segments": [seg("texte")], "mise_en_page": {"alignement": "droite", "police": "Arial"}}]}).json()["adresse"]
    cle = adresse.split("/")[1]
    blocs = editeur(c, cle)["blocs"]
    blocs.append({"type": "paragraphe", "id": None, "origine": blocs[0]["id"], "segments": [seg("ajout")]})
    r = c.put(f"/api/lot/{cle}/document/0", json={"blocs": blocs})
    assert r.json()["blocs"][1]["mise_en_page"]["alignement"] == "droite"


def test_coller_vide_ou_invalide(app):
    c = client(app)
    r = c.post("/api/coller", json={"blocs": [{"segments": [seg("  ")]}]})
    assert r.status_code == 422 and "vide" in r.json()["erreur"]
    assert c.post("/api/coller", json={"blocs": []}).status_code == 422
    assert c.post("/api/coller", json={"blocs": [{"segments": [], "mise_en_page": {"taille": 9999}}]}).status_code == 422


def test_coller_document_prive(app):
    adresse = client(app).post("/api/coller", json={"blocs": [{"segments": [seg("texte")]}]}).json()["adresse"]
    assert client(app, "autre").get(f"/{adresse}").status_code == 404


def test_mise_en_page_d_un_document_depose(app):
    """« Copier pour Word » fonctionne aussi depuis un fichier déposé : police du style à défaut."""
    from tests.fabrique import compte_rendu

    c = client(app)
    r = c.post("/formater", files={"fichiers": ("cr.docx", compte_rendu(), DOCX)})
    cle = r.headers["location"].rsplit("/", 1)[1]
    blocs = editeur(c, cle)["blocs"]
    titre = next(b for b in blocs if b["type"] == "paragraphe" and "COMPTE-RENDU" in textes([b])[0])
    assert titre["mise_en_page"]["alignement"] == "centre"
    assert next(b for b in blocs if b["type"] == "tableau")["mise_en_page"] == {"bordures": False}


# ---- Accès direct (proxy nginx) -------------------------------------------------------

NGINX = "172.30.33.5"


@pytest.fixture
def app_directe(settings):
    settings.acces_direct = True
    settings.utilisateurs = {"secretaire": "un-mot-de-passe-solide"}
    return creer_app(settings)


def navigateur(app, **en_tetes):
    c = TestClient(app, client=(NGINX, 50000), follow_redirects=False)
    c.headers.update({"Accept": "text/html", "X-Real-IP": "203.0.113.7", **en_tetes})
    return c


def se_connecter(c, mot_de_passe="un-mot-de-passe-solide", suite="/"):
    return c.post("/connexion", data={"identifiant": "Secretaire", "mot_de_passe": mot_de_passe, "suite": suite})


def test_acces_direct_desactive_par_defaut(app):
    assert navigateur(app).get("/").status_code == 403


def test_acces_direct_exige_connexion(app_directe):
    c = navigateur(app_directe)
    r = c.get("/vocabulaire")
    assert r.status_code == 303 and r.headers["location"] == "/connexion?suite=/vocabulaire"
    assert c.get("/api/vocabulaire", headers={"Accept": "application/json"}).status_code == 401
    assert c.get("/connexion").status_code == 200
    assert c.get("/static/style.css").status_code == 200


def test_connexion_et_deconnexion(app_directe):
    c = navigateur(app_directe)
    r = se_connecter(c, suite="/vocabulaire")
    assert r.status_code == 303 and r.headers["location"] == "/vocabulaire"
    cookie = r.headers["set-cookie"]
    assert "HttpOnly" in cookie and "SameSite=lax" in cookie
    page = c.get("/vocabulaire")
    assert page.status_code == 200 and "Se déconnecter (secretaire)" in page.text
    # Les en-têtes Ingress envoyés par le client sont ignorés
    assert '<base href="/">' in c.get("/", headers={"X-Ingress-Path": "//pirate.example"}).text
    # L'historique enregistre l'identifiant de connexion
    c.post("/api/vocabulaire/mot", json={"mot": "apixaban3"})
    assert "secretaire" in c.get("/historique").text

    c.post("/deconnexion")
    assert c.get("/vocabulaire").status_code == 303


def test_cookie_securise_derriere_https(app_directe):
    r = se_connecter(navigateur(app_directe, **{"X-Forwarded-Proto": "https"}))
    assert "Secure" in r.headers["set-cookie"]


def test_mauvais_mot_de_passe_et_blocage(app_directe):
    c = navigateur(app_directe)
    for _ in range(5):
        r = se_connecter(c, mot_de_passe="faux")
        assert r.status_code == 401 and "incorrect" in r.text
    r = se_connecter(c)  # même le bon mot de passe est refusé pendant le blocage
    assert r.status_code == 401 and "Trop de tentatives" in r.text
    # Une autre adresse n'est pas bloquée
    assert se_connecter(navigateur(app_directe, **{"X-Real-IP": "198.51.100.2"})).status_code == 303


def test_session_falsifiee_ou_mot_de_passe_change(app_directe, settings):
    c = navigateur(app_directe)
    jeton = se_connecter(c).cookies.get("docformatter_session")
    assert jeton
    falsifie = base64.urlsafe_b64encode(b"admin|9999999999|00").decode()
    c.cookies.set("docformatter_session", falsifie)
    assert c.get("/vocabulaire").status_code == 303

    # Un changement de mot de passe invalide les sessions existantes
    settings.utilisateurs = {"secretaire": "nouveau-mot-de-passe-solide"}
    autre = navigateur(creer_app(settings))
    autre.cookies.set("docformatter_session", jeton)
    assert autre.get("/vocabulaire").status_code == 303


def test_redirection_externe_refusee(app_directe):
    r = se_connecter(navigateur(app_directe), suite="//pirate.example/")
    assert r.headers["location"] == "/"


def test_ingress_toujours_possible_en_acces_direct(app_directe):
    assert client(app_directe).get("/vocabulaire").status_code == 200


def test_options_utilisateurs(tmp_path, monkeypatch):
    from app.settings import charger_settings

    options = tmp_path / "options.json"
    options.write_text(json.dumps({"acces_direct": True, "utilisateurs": [
        {"nom": "Secretaire ", "mot_de_passe": "un-mot-de-passe-solide"},
        {"nom": "court", "mot_de_passe": "trop-court"[:9]},
    ]}))
    monkeypatch.setenv("DOCFORMATTER_OPTIONS", str(options))
    monkeypatch.setenv("DOCFORMATTER_CONFIG", str(tmp_path / "config"))
    monkeypatch.setenv("DOCFORMATTER_DATA", str(tmp_path / "data"))
    s = charger_settings()
    assert s.acces_direct and s.utilisateurs == {"secretaire": "un-mot-de-passe-solide"}
