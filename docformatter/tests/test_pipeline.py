import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from app.pipeline import finaliser_retouche, formater, traiter_texte
from app.pipeline.model import Bloc, Format, Inconnu, TexteStyle
from app.pipeline.reader import lire_docx
from app.pipeline.spelling import Correcteur
from app.pipeline.writer import ecrire_docx
from app.retouche import Edition, depuis_editeur, remplacer_mots, vers_editeur
from app.rules import Regles, lire_yaml
from app.store import RegleInvalide
from fabrique import compte_rendu

NBSP = "\u00a0"
DEFAUTS = Path(__file__).resolve().parents[1] / "defaults" / "regles.yaml"


@pytest.fixture
def regles() -> Regles:
    return lire_yaml(DEFAUTS.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def original() -> bytes:
    return compte_rendu()


def _xml(docx: bytes, partie: str) -> bytes:
    return etree.tostring(etree.fromstring(zipfile.ZipFile(io.BytesIO(docx)).read(partie)), method="c14n")


def _paragraphes(docx: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(docx)).paragraphs]


def _corps(docx: bytes):
    return Document(io.BytesIO(docx)).element.body


def _compter(corps, balise: str) -> int:
    return len(corps.findall(".//" + balise, corps.nsmap))


# ---- Règles par défaut ----------------------------------------------------------------

def test_exemples_des_regles(regles):
    attendus = {
        "Points de suspension": "Douleur… puis amélioration",
        "Espace après une virgule": "HTA, diabète, créatinine 1,2 mg/dL",
        "Millilitres en mL": f"500{NBSP}mL de sérum physiologique",
        "Espacer les unités": f"Kardégic 75{NBSP}mg, poids 72{NBSP}kg, SpO2 98{NBSP}%",
        "Espaces avant : ; ! ?": f"TA{NBSP}: 13/8{NBSP}; FC{NBSP}: 72 bpm. RDV à 10:30",
        "Espace après les deux-points": "Motif: douleur thoracique",
    }
    par_nom = {r.nom: r for r in regles.remplacements}
    for nom, attendu in attendus.items():
        assert par_nom[nom].appliquer(par_nom[nom].exemple)[0] == attendu, nom
    # Chaque règle par défaut a une description et un exemple qui la déclenche (page « Mise en forme »)
    for r in regles.remplacements:
        assert r.description and r.appliquer(r.exemple)[1], r.nom


# ---- Mise en page conservée -----------------------------------------------------------

def test_entetes_pieds_sections_conserves(original, regles):
    sortie = formater(original, regles, None).docx
    for partie in zipfile.ZipFile(io.BytesIO(original)).namelist():
        if partie.endswith(".xml") and partie != "word/document.xml" and "Content_Types" not in partie:
            assert _xml(original, partie) == _xml(sortie, partie), partie
    avant, apres = _corps(original), _corps(sortie)
    for balise in ("w:sectPr", "w:cols", "w:titlePg", "w:headerReference", "w:footerReference",
                   "w:hyperlink", "w:bookmarkEnd", "w:tbl", "w:tab"):
        assert _compter(avant, balise) == _compter(apres, balise), balise
    assert len(_paragraphes(sortie)) == len(_paragraphes(original))


def test_paragraphes_inchanges_identiques(original, regles):
    """Un paragraphe sans correction n'est pas réécrit : son XML reste strictement identique."""
    sortie = formater(original, regles, None).docx
    avant, apres = list(_corps(original)), list(_corps(sortie))
    inchanges = [i for i, p in enumerate(Document(io.BytesIO(original)).paragraphs)
                 if p.text in ("Ville, le lundi 5 octobre 2026", "Chère Consœur, Cher Confrère,")]
    assert inchanges
    for i in inchanges:
        assert etree.tostring(avant[i], method="c14n") == etree.tostring(apres[i], method="c14n")


def test_corrections_sur_place_avec_mise_en_forme(original, regles):
    res = formater(original, regles, None)
    textes = _paragraphes(res.docx)
    assert f"Antécédents{NBSP}: HTA, diabète type 2" in textes  # rubrique renommée + typographie + correction
    assert "Douleur thoracique depuis 3 jours…" in textes
    assert f"Poids{NBSP}: 72{NBSP}kg" in textes
    assert "-Suivi pulmonaire Dr Martin" in textes  # majuscule après le tiret de liste

    doc = Document(io.BytesIO(res.docx))
    atcd = next(p for p in doc.paragraphs if p.text.startswith("Antécédents"))
    assert atcd.runs[0].text == "Antécédents" and atcd.runs[0].bold and atcd.runs[0].underline
    assert not atcd.runs[1].bold
    motif = next(p for p in doc.paragraphs if p.text.startswith("Motif"))
    assert motif.runs[0].text == "Motif de consultation" and motif.runs[0].bold
    assert motif.text == f"Motif de consultation{NBSP}:"


def test_alignement_et_paragraphes_proteges(original, regles):
    textes = _paragraphes(formater(original, regles, None).docx)
    # Tabulation et espaces d'alignement de la signature conservés
    assert "\t                    Docteur Martin" in textes
    # Paragraphe contenant un lien : intact (pas d'espace insécable ajoutée avant « : »)
    assert "Via la plateforme de télé-expertise : www.exemple-fictif.fr" in textes


def test_tableau_et_section_a_deux_colonnes(original, regles):
    sortie = formater(original, regles, None).docx
    doc = Document(io.BytesIO(sortie))
    assert doc.tables[0].cell(0, 1).text == f"96{NBSP}µmol/L"
    assert f"Hb (g/dL){NBSP}: 13,4" in _paragraphes(sortie)
    assert len(doc.sections) == 3 and doc.sections[0].different_first_page_header_footer


def test_rubrique_deja_correcte_non_modifiee(regles):
    res = traiter_texte("MOTIF DE CONSULTATION :\nAntecedents : HTA", regles, None)
    assert [b.texte for b in res.blocs] == [f"MOTIF DE CONSULTATION{NBSP}:", f"Antecedents{NBSP}: HTA"]


def test_rubrique_sans_separateur_non_modifiee(regles):
    res = traiter_texte("Traitement par IPP débuté hier", regles, None)
    assert res.blocs[0].texte == "Traitement par IPP débuté hier"


def test_corrections_conservent_la_casse(regles):
    res = traiter_texte("Hypertention connue, OEDEME des membres inférieurs, tachychardie", regles, None)
    assert res.blocs[0].texte == "Hypertension connue, ŒDÈME des membres inférieurs, tachycardie"


def test_correction_apres_tiret_et_deux_points_colles(regles):
    res = traiter_texte("-kardegic 75mg\nMotif:douleur", regles, None)
    assert [b.texte for b in res.blocs] == [f"-Kardégic 75{NBSP}mg", f"Motif de consultation{NBSP}: douleur"]


def test_majuscule_debut_respecte_les_unites(regles):
    res = traiter_texte("pH à 7,32\nmise sous antibiotiques", regles, None)
    assert [b.texte for b in res.blocs] == ["pH à 7,32", "Mise sous antibiotiques"]


# ---- Texte avec mise en forme ---------------------------------------------------------

def test_texte_style_garde_la_mise_en_forme():
    bloc = Bloc("paragraphe", texte="ATCD : HTA", formats=[Format(0, 4, gras=True), Format(4, 10)])
    texte = TexteStyle(bloc)
    texte.remplacer(0, 4, "Antécédents")
    texte.remplacer(len("Antécédents"), len("Antécédents") + 1, NBSP)
    texte.appliquer(bloc)
    assert bloc.texte == f"Antécédents{NBSP}: HTA"
    assert [(f.debut, f.fin, f.gras) for f in bloc.formats] == [(0, 11, True), (11, 17, False)]


def test_mot_inconnu_a_cheval_sur_deux_mises_en_forme(original):
    blocs = lire_docx(original)
    p = next(b for b in blocs if b.texte.startswith("Motif"))
    p.inconnus = [Inconnu(18, 22, "ion ", ["x"])]
    doc = Document(io.BytesIO(ecrire_docx(original, blocs, commentaires=True)))
    motif = next(q for q in doc.paragraphs if q.text.startswith("Motif"))
    assert [(r.text, bool(r.bold)) for r in motif.runs if r.text] == [
        ("Motif de consultat", True), ("ion", True), (" ", False), (":", False)]
    assert len(list(doc.comments)) == 1


@pytest.mark.skipif(not Correcteur([]).disponible, reason="hunspell non installé")
def test_orthographe_sur_place(original, regles):
    res = formater(original, regles, Correcteur(regles.dictionnaire))
    doc = Document(io.BytesIO(res.docx))
    assert all(c.author == "DocFormatter" for c in doc.comments)
    res2 = traiter_texte("Patient présentant une dispnée.\n-suivi pulmonaire", regles, Correcteur([]))
    assert res2.blocs[1].inconnus == []  # « -Suivi » : le tiret de liste ne fait pas partie du mot
    assert [i.mot for i in res2.blocs[0].inconnus] == ["dispnée"]
    assert "dyspnée" in res2.blocs[0].inconnus[0].suggestions


# ---- Retouche sur place ---------------------------------------------------------------

def _edition(donnees):
    return Edition.model_validate({"blocs": donnees})


def test_retouche_modifier_ajouter_supprimer_deplacer(original, regles):
    blocs = formater(original, regles, None).blocs
    donnees = vers_editeur(blocs)
    i4 = next(i for i, b in enumerate(donnees) if b.get("id") == "p4")
    donnees[i4]["segments"] = [{"texte": "Cher ", "gras": False}, {"texte": "Confrère", "gras": True}]
    # Nouveau paragraphe après « Cher Confrère », sur le modèle de p4
    donnees.insert(i4 + 1, {"type": "paragraphe", "id": None, "origine": "p4",
                            "segments": [{"texte": "Ligne ajoutée", "italique": True}]})
    # Suppression de « Poids » et échange des deux premières lignes de biologie
    donnees = [b for b in donnees if b.get("id") != "p9"]
    i12 = next(i for i, b in enumerate(donnees) if b.get("id") == "p12")
    donnees[i12], donnees[i12 + 1] = donnees[i12 + 1], donnees[i12]

    retouches = depuis_editeur(_edition(donnees), blocs)
    sortie = finaliser_retouche(original, retouches, regles, None).docx
    textes = _paragraphes(sortie)
    assert textes[textes.index("Cher Confrère") + 1] == "Ligne ajoutée"
    assert not any(t.startswith("Poids") for t in textes)
    assert textes.index(f"Plaquettes (Giga/L){NBSP}: 294") < textes.index(f"Hb (g/dL){NBSP}: 13,4")

    doc = Document(io.BytesIO(sortie))
    cher = next(p for p in doc.paragraphs if p.text == "Cher Confrère")
    assert [(r.text, bool(r.bold)) for r in cher.runs] == [("Cher ", False), ("Confrère", True)]
    assert next(p for p in doc.paragraphs if p.text == "Ligne ajoutée").runs[0].italic
    # Mise en page conservée : sections, en-têtes, lien, signet
    assert len(doc.sections) == 3 and doc.sections[0].different_first_page_header_footer
    assert _xml(original, "word/header2.xml") == _xml(sortie, "word/header2.xml")
    assert _compter(_corps(sortie), "w:hyperlink") == 1 and _compter(_corps(sortie), "w:bookmarkEnd") == 1


def test_retouche_sans_modification_identique(original, regles):
    blocs = formater(original, regles, None).blocs
    retouches = depuis_editeur(_edition(vers_editeur(blocs)), blocs)
    avant = _corps(ecrire_docx(original, blocs, False))
    apres = _corps(ecrire_docx(original, retouches, False))
    assert etree.tostring(avant, method="c14n") == etree.tostring(apres, method="c14n")


def test_retouche_protections(original, regles):
    blocs = formater(original, regles, None).blocs
    donnees = vers_editeur(blocs)
    with pytest.raises(RegleInvalide, match="protégé"):
        depuis_editeur(_edition([b for b in donnees if not b.get("fin_section")]), blocs)
    with pytest.raises(RegleInvalide, match="tableau"):
        depuis_editeur(_edition([b for b in donnees if b["type"] != "tableau"]), blocs)
    with pytest.raises(RegleInvalide, match="rechargez"):
        depuis_editeur(_edition(donnees + [{"type": "paragraphe", "id": "p999", "segments": []}]), blocs)
    # Un paragraphe protégé modifié dans l'éditeur reste tel quel
    next(b for b in donnees if b.get("protege"))["segments"] = [{"texte": "piraté"}]
    retouches = depuis_editeur(_edition(donnees), blocs)
    assert "Via la plateforme de télé-expertise : www.exemple-fictif.fr" in _paragraphes(
        finaliser_retouche(original, retouches, regles, None).docx)


def test_remplacer_mots_garde_la_mise_en_forme(original, regles):
    blocs = remplacer_mots(formater(original, regles, None).blocs, {"consultation": "visite"})
    motif = next(b for b in blocs if b.texte.startswith("Motif"))
    assert motif.texte.startswith("Motif de visite") and motif.formats[0].gras


def test_retour_a_la_ligne(original, regles):
    blocs = formater(original, regles, None).blocs
    donnees = vers_editeur(blocs)
    next(b for b in donnees if b.get("id") == "p6")["segments"] = [{"texte": "Première ligne\nDeuxième ligne"}]
    sortie = finaliser_retouche(original, depuis_editeur(_edition(donnees), blocs), regles, None).docx
    assert "Première ligne\nDeuxième ligne" in _paragraphes(sortie)
