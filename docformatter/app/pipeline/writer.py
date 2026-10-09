"""Écriture du .docx final à partir d'un modèle Word (styles, en-tête, pied de page)."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor

from .model import Bloc

AUTEUR = "DocFormatter"


def modele_par_defaut() -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    titre = doc.styles["Heading 1"]
    titre.font.name = "Calibri"
    titre.font.size = Pt(13)
    titre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    titre.paragraph_format.space_before = Pt(14)
    titre.paragraph_format.space_after = Pt(4)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _vider_corps(doc) -> None:
    corps = doc.element.body
    for enfant in list(corps):
        if not enfant.tag.endswith("}sectPr"):
            corps.remove(enfant)


def _style(doc, *noms: str):
    for nom in noms:
        try:
            return doc.styles[nom]
        except KeyError:
            continue
    return None


def _ecrire_texte(doc, paragraphe, bloc: Bloc, commentaires: bool) -> None:
    """Ajoute le texte en surlignant les mots inconnus (+ commentaire avec suggestions)."""
    pos = 0
    for inconnu in sorted(bloc.inconnus, key=lambda x: x.debut):
        if inconnu.debut > pos:
            paragraphe.add_run(bloc.texte[pos : inconnu.debut])
        run = paragraphe.add_run(bloc.texte[inconnu.debut : inconnu.fin])
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if commentaires:
            texte = (
                "Mot inconnu. Suggestions : " + ", ".join(inconnu.suggestions)
                if inconnu.suggestions
                else "Mot inconnu, aucune suggestion."
            )
            doc.add_comment(run, text=texte, author=AUTEUR, initials="DF")
        pos = inconnu.fin
    if pos < len(bloc.texte):
        paragraphe.add_run(bloc.texte[pos:])


def ecrire_docx(blocs: list[Bloc], modele: bytes | None, commentaires: bool) -> bytes:
    doc = Document(io.BytesIO(modele or modele_par_defaut()))
    _vider_corps(doc)

    style_titre = _style(doc, "Heading 1", "Titre 1")
    style_puce = _style(doc, "List Bullet", "Liste à puces")
    style_num = _style(doc, "List Number", "Liste à numéros")
    style_tableau = _style(doc, "Table Grid", "Grille du tableau")

    for bloc in blocs:
        if bloc.type == "titre":
            p = doc.add_paragraph(style=style_titre)
            _ecrire_texte(doc, p, bloc, commentaires)
        elif bloc.type == "liste":
            style = style_num if bloc.numerote else style_puce
            p = doc.add_paragraph(style=style)
            if style is None:
                p.add_run("• ")
            _ecrire_texte(doc, p, bloc, commentaires)
        elif bloc.type == "tableau":
            if not bloc.lignes:
                continue
            nb_col = max(len(ligne) for ligne in bloc.lignes)
            table = doc.add_table(rows=len(bloc.lignes), cols=nb_col)
            if style_tableau is not None:
                table.style = style_tableau
            for i, ligne in enumerate(bloc.lignes):
                for j, cellule in enumerate(ligne):
                    p = table.cell(i, j).paragraphs[0]
                    _ecrire_texte(doc, p, cellule, commentaires)
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            _ecrire_texte(doc, p, bloc, commentaires)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def valider_modele(data: bytes) -> None:
    """Lève une exception si le fichier n'est pas un .docx utilisable comme modèle."""
    ecrire_docx([Bloc("titre", texte="Test"), Bloc("paragraphe", texte="Test")], data, False)
