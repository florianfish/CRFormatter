"""Compte rendu collé depuis Word.

Le navigateur lit le contenu collé (texte, gras / italique / souligné, mise en page de chaque
paragraphe, tableaux) ; il est converti ici en .docx pour suivre exactement le même traitement
qu'un fichier déposé (règles, orthographe, retouche). Le résultat repart dans Word par
« Copier pour Word » ; en-têtes et pieds de page n'ont jamais quitté le document d'origine.
"""

from __future__ import annotations

import io
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, Field

from .retouche import Segment

ALIGNEMENTS = {"centre": WD_ALIGN_PARAGRAPH.CENTER, "droite": WD_ALIGN_PARAGRAPH.RIGHT,
               "justifie": WD_ALIGN_PARAGRAPH.JUSTIFY}


class MiseEnPage(BaseModel):
    """Mesures en points."""

    alignement: Literal["centre", "droite", "justifie"] | None = None
    retrait_gauche: float | None = Field(None, ge=-500, le=1000)
    retrait_premiere_ligne: float | None = Field(None, ge=-500, le=500)
    espace_avant: float | None = Field(None, ge=0, le=500)
    espace_apres: float | None = Field(None, ge=0, le=500)
    interligne: float | None = Field(None, ge=0.5, le=5)
    police: str | None = Field(None, max_length=100)
    taille: float | None = Field(None, ge=1, le=200)


class ParagrapheColle(BaseModel):
    segments: list[Segment] = Field(default_factory=list, max_length=2_000)
    mise_en_page: MiseEnPage = Field(default_factory=MiseEnPage)


class BlocColle(ParagrapheColle):
    type: Literal["paragraphe", "tableau"] = "paragraphe"
    bordures: bool = False
    lignes: list[list[list[ParagrapheColle]]] = Field(default_factory=list, max_length=500)


class Collage(BaseModel):
    blocs: list[BlocColle] = Field(min_length=1, max_length=5_000)

    def vide(self) -> bool:
        paragraphes = [p for b in self.blocs for p in ([b] if b.type == "paragraphe" else
                                                       [p for ligne in b.lignes for c in ligne for p in c])]
        return not any(s.texte.strip() for p in paragraphes for s in p.segments)


def _police_du_paragraphe(paragraphe, m: MiseEnPage) -> None:
    """Police de la marque de paragraphe : garde la hauteur d'une ligne vide."""
    rpr = OxmlElement("w:rPr")
    if m.police:
        polices = OxmlElement("w:rFonts")
        for attribut in ("w:ascii", "w:hAnsi", "w:cs"):
            polices.set(qn(attribut), m.police)
        rpr.append(polices)
    if m.taille:
        taille = OxmlElement("w:sz")
        taille.set(qn("w:val"), str(round(m.taille * 2)))
        rpr.append(taille)
    if len(rpr):
        paragraphe._p.get_or_add_pPr().append(rpr)


def _remplir(paragraphe, saisi: ParagrapheColle) -> None:
    m, f = saisi.mise_en_page, paragraphe.paragraph_format
    if m.alignement:
        f.alignment = ALIGNEMENTS[m.alignement]
    if m.retrait_gauche:
        f.left_indent = Pt(m.retrait_gauche)
    if m.retrait_premiere_ligne:
        f.first_line_indent = Pt(m.retrait_premiere_ligne)
    # Toujours explicites : sans eux, l'espacement du style par défaut s'appliquerait
    f.space_before, f.space_after = Pt(m.espace_avant or 0), Pt(m.espace_apres or 0)
    if m.interligne:
        f.line_spacing = m.interligne
    for seg in saisi.segments:
        texte = seg.texte.replace("\r", "")
        if not texte:
            continue
        run = paragraphe.add_run(texte)  # « \t » et « \n » deviennent tabulation et retour à la ligne
        run.bold, run.italic, run.underline = seg.gras or None, seg.italique or None, seg.souligne or None
        if m.police:
            run.font.name = m.police
        if m.taille:
            run.font.size = Pt(m.taille)
    _police_du_paragraphe(paragraphe, m)


def docx_depuis_collage(collage: Collage) -> bytes:
    doc = Document()
    for bloc in collage.blocs:
        if bloc.type == "paragraphe":
            _remplir(doc.add_paragraph(), bloc)
            continue
        lignes = [ligne for ligne in bloc.lignes if ligne]
        if not lignes:
            continue
        table = doc.add_table(rows=len(lignes), cols=max(len(ligne) for ligne in lignes))
        if bloc.bordures:
            table.style = "Table Grid"
        for r, ligne in enumerate(lignes):
            for c, paragraphes in enumerate(ligne):
                cellule = table.cell(r, c)
                for k, p in enumerate(paragraphes or [ParagrapheColle()]):
                    _remplir(cellule.paragraphs[0] if k == 0 else cellule.add_paragraph(), p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
