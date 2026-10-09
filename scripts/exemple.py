"""Génère un compte rendu fictif volontairement mal rédigé, présenté comme un courrier
hospitalier (en-tête de première page, en-tête des pages suivantes, pied de page).

- exemple.docx : à ouvrir dans Word (ou LibreOffice) pour copier le texte et le coller dans l'outil ;
- exemple.html : le même texte tel que Word le place dans le presse-papiers (utilisé par `make e2e`).
"""

import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt

sortie = Path(sys.argv[1] if len(sys.argv) > 1 else "exemple.docx")
sortie.parent.mkdir(parents=True, exist_ok=True)

d = Document()
section = d.sections[0]
section.left_margin = section.right_margin = Cm(1.5)
section.different_first_page_header_footer = True
section.first_page_header.paragraphs[0].text = "HÔPITAL FICTIF — Service de médecine interne — Dr Martin"
section.header.paragraphs[0].text = "DUPONT Jean, né le 01/01/1950, page "
section.header.paragraphs[0]._p.append(parse_xml(
    f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
for pied in (section.footer, section.first_page_footer):
    pied.paragraphs[0].text = "Hôpital fictif – 1 rue de l'Exemple – 00000 VILLE"


def rubrique(libelle, suite=""):
    p = d.add_paragraph()
    r = p.add_run(libelle)
    r.bold = r.underline = True
    if suite:
        p.add_run(suite)


titre = d.add_paragraph()
titre.alignment = 1
r = titre.add_run("COMPTE-RENDU DE CONSULTATION")
r.bold, r.font.size = True, Pt(13)
d.add_paragraph("")
rubrique("Motif", ":douleur thoracique depuis 3 jours...")
rubrique("ATCD", " : HTA , diabéte type 2,dyslipidémie")
rubrique("ttt habituel", ":")
d.add_paragraph("-kardegic 75mg\t1/j")
d.add_paragraph("- metformine 1000mg x2/j")
rubrique("Examen clinique", " :")
p = d.add_paragraph("patient eupnéique,apyrétique.")
p.add_run().add_break()  # retour à la ligne manuel (Maj+Entrée)
p.add_run("TA: 14/9 ; FC: 88bpm. Douleur \"en étau\", pas d'oedeme des MI. dispnée d'effort.")
rubrique("Biologie", " :")
t = d.add_table(rows=2, cols=2)
for i, (nom, valeur) in enumerate([("Troponine", "12ng/L"), ("Créat", "80µmol/L")]):
    t.cell(i, 0).text, t.cell(i, 1).text = nom, valeur
rubrique("CAT", " : coroscanner a programmer, revoir en consultation dans 1 mois")
d.add_paragraph("")
d.add_paragraph("\t\t\t\t\t\tDocteur Martin")
d.save(sortie)

# Corps du compte rendu tel que Word le copie (feuille de style MsoNormal, rubriques en gras souligné)
RUBRIQUE = "<b><u>{}</u></b>"
html = f"""<html xmlns:o="urn:schemas-microsoft-com:office:office"><head><meta charset="utf-8"><style>
p.MsoNormal {{margin:0cm; margin-bottom:8.0pt; line-height:107%; font-size:11.0pt; font-family:"Calibri",sans-serif;}}
</style></head><body lang=FR><div class=WordSection1>
<p class=MsoNormal align=center style='text-align:center'><b><span style='font-size:13.0pt'>COMPTE-RENDU DE CONSULTATION</span></b></p>
<p class=MsoNormal><o:p>&nbsp;</o:p></p>
<p class=MsoNormal>{RUBRIQUE.format("Motif")}:douleur thoracique depuis 3 jours...</p>
<p class=MsoNormal>{RUBRIQUE.format("ATCD")} : HTA , diabéte type 2,dyslipidémie</p>
<p class=MsoNormal>{RUBRIQUE.format("ttt habituel")}:</p>
<p class=MsoNormal>-kardegic 75mg<span style='mso-tab-count:1'>      </span>1/j</p>
<p class=MsoNormal>- metformine 1000mg x2/j</p>
<p class=MsoNormal>{RUBRIQUE.format("Examen clinique")} :</p>
<p class=MsoNormal>patient eupnéique,apyrétique.<br>TA: 14/9 ; FC: 88bpm. Douleur "en étau", pas d'oedeme
des MI. dispnée d'effort.</p>
<p class=MsoNormal>{RUBRIQUE.format("Biologie")} :</p>
<table class=MsoNormalTable border=0 cellspacing=0 cellpadding=0 style='border-collapse:collapse'>
<tr><td><p class=MsoNormal>Troponine</p></td><td><p class=MsoNormal>12ng/L</p></td></tr>
<tr><td><p class=MsoNormal>Créat</p></td><td><p class=MsoNormal>80µmol/L</p></td></tr>
</table>
<p class=MsoNormal>{RUBRIQUE.format("CAT")} : coroscanner a programmer, revoir en consultation dans 1 mois</p>
<p class=MsoNormal><o:p>&nbsp;</o:p></p>
<p class=MsoNormal>{"<span style='mso-tab-count:6'>      </span>"}Docteur Martin</p>
</div></body></html>"""
sortie.with_suffix(".html").write_text(html, encoding="utf-8")
print(f"Exemple écrit : {sortie} et {sortie.with_suffix('.html')}")
