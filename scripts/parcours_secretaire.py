"""Parcours de la secrétaire dans un vrai navigateur (Playwright + Google Chrome).

À lancer sur une instance de développement : `make e2e`.
Usage : parcours_secretaire.py URL EXEMPLE.docx DOSSIER_CAPTURES
"""
import sys
from playwright.sync_api import sync_playwright, expect

URL, EXEMPLE, SP = sys.argv[1], sys.argv[2], sys.argv[3]
erreurs = []

with sync_playwright() as p:
    nav = p.chromium.launch(channel="chrome", args=["--no-sandbox"])
    page = nav.new_page(viewport={"width": 1100, "height": 900})
    page.on("pageerror", lambda e: erreurs.append(str(e)))
    page.on("console", lambda m: m.type == "error" and erreurs.append(m.text))

    # 1. Dépôt du document
    page.goto(URL)
    page.set_input_files("#fichiers", EXEMPLE)
    page.click("#envoyer")
    page.wait_for_url("**/lot/**")
    expect(page.locator(".verification")).to_contain_text("dispnée")
    expect(page.locator("#notification")).to_be_hidden()
    print("1. document formaté, mots à vérifier affichés")

    # 2. Taper un remplacement coche « Remplacer par » ; « Créat » est déclaré correct
    mot = page.locator("fieldset.mot", has_text="dispnée")
    mot.locator("input[type=text]").fill("dyspnée")
    assert mot.locator("input[value=remplacer]").is_checked()
    page.locator("fieldset.mot", has_text="Créat").get_by_label("Le mot est correct").check()
    page.click("text=Enregistrer mes choix et reformater")
    expect(page.locator(".message.succes")).to_contain_text("C'est noté")
    expect(page.locator(".verification")).to_have_count(0)
    expect(page.locator(".apercu")).to_contain_text("dyspnée")
    page.screenshot(path=f"{SP}/apres-choix.png", full_page=True)
    print("2. choix enregistrés, document reformaté sans mot à vérifier")

    # 3. Bouton « Annuler » du message : les mots reviennent
    page.click("#annuler-decisions")
    page.wait_for_load_state("networkidle")
    expect(page.locator(".verification")).to_contain_text("dispnée")
    print("3. annulation depuis le résultat : mots à vérifier de retour")

    # 4. Vocabulaire : ajout d'un remplacement puis annulation via la notification
    page.goto(URL + "vocabulaire")
    page.fill("#form-remplacement [name=texte]", "ttt")
    page.fill("#form-remplacement [name=par]", "traitement")
    page.click("#form-remplacement button")
    expect(page.locator("#liste-remplacements")).to_contain_text("traitement")
    expect(page.locator("#notification")).to_contain_text("Remplacement ajouté")
    page.screenshot(path=f"{SP}/notification.png")
    page.click("#notification button")
    page.wait_for_load_state("networkidle")
    expect(page.locator("#liste-remplacements")).not_to_contain_text("ttt")
    print("4. remplacement ajouté puis annulé depuis la notification")

    # 5. Erreur compréhensible
    page.fill("#form-mot [name=mot]", "deux mots")
    page.click("#form-mot button")
    expect(page.locator("#notification")).to_contain_text("Un seul mot")
    print("5. message d'erreur clair")

    # 6. Rubrique : ajout d'une écriture reconnue
    rubrique = page.locator(".rubrique", has_text="Antécédents chirurgicaux")
    rubrique.locator("input").fill("atcd chirurgie")
    rubrique.locator("button[type=submit]").click()
    expect(page.locator(".rubrique", has_text="Antécédents chirurgicaux")).to_contain_text("atcd chirurgie")
    print("6. écriture de rubrique ajoutée")

    # 7. Mise en forme : interrupteur
    page.goto(URL + "mise-en-forme")
    page.locator(".interrupteur", has_text="Guillemets français").click()
    expect(page.locator("#notification")).to_contain_text("désactivé")
    page.reload()
    assert not page.locator(".interrupteur", has_text="Guillemets français").locator("input").is_checked()
    print("7. interrupteur désactivé et conservé après rechargement")

    # 8. Historique : retour à la version initiale
    page.goto(URL + "historique")
    page.screenshot(path=f"{SP}/historique.png", full_page=True)
    page.on("dialog", lambda d: d.accept())
    page.locator(".historique li").last.locator("button").click()
    page.wait_for_load_state("networkidle")
    expect(page.locator(".historique li").first).to_contain_text("Retour à la version du")
    page.goto(URL + "mise-en-forme")
    assert page.locator(".interrupteur", has_text="Guillemets français").locator("input").is_checked()
    print("8. retour à l'état initial depuis l'historique")

    # 9. Mode expert accessible
    assert page.goto(URL + "expert").status == 200
    expect(page.locator("#liste-regles .regle").first).to_be_visible()
    print("9. mode expert accessible")

    # 10. Retouche manuelle du document (sur place : en-têtes et mise en page conservés)
    page.goto(URL)
    page.set_input_files("#fichiers", EXEMPLE)
    page.click("#envoyer")
    page.wait_for_url("**/lot/**")
    page.click("text=Retoucher")
    page.wait_for_selector("#feuille .bloc")
    expect(page.locator("#feuille mark.inconnu").first).to_be_visible()

    # Modifier un paragraphe : curseur à la fin, saisie, puis Entrée pour en créer un nouveau
    paragraphe = page.locator("#feuille p.bloc", has_text="douleur thoracique")
    paragraphe.click()
    page.keyboard.press("End")
    page.keyboard.type(" (retouche)")
    page.keyboard.press("Enter")
    page.keyboard.type("Nouveau paragraphe saisi")
    expect(page.locator("#etat")).to_have_text("Modifications non enregistrées")

    # Annuler : la frappe puis la création du paragraphe
    page.click('[data-action="annuler"]')
    page.click('[data-action="annuler"]')
    expect(page.locator("#feuille p.bloc", has_text="Nouveau paragraphe saisi")).to_have_count(0)
    paragraphe.click()
    page.keyboard.press("End")
    page.keyboard.press("Enter")
    page.keyboard.type("Nouveau paragraphe saisi")

    # Retour arrière en début de paragraphe : fusion avec le précédent
    page.locator("#feuille p.bloc", has_text="Nouveau paragraphe saisi").click()
    page.keyboard.press("Home")
    page.keyboard.press("Backspace")
    expect(page.locator("#feuille p.bloc", has_text="(retouche)Nouveau paragraphe saisi")).to_be_visible()

    # Remplacer un mot inconnu dans ce document depuis le panneau
    mot = page.locator(".liste-mots li", has_text="dispnée")
    mot.locator("input[type=text]").fill("dyspnée")
    mot.get_by_role("button", name="Remplacer").click()
    expect(page.locator("#feuille")).to_contain_text("dyspnée d'effort")

    # Gras avec le bouton, italique avec Ctrl+I, sur le dernier mot sélectionné au clavier
    examen = page.locator("#feuille p.bloc", has_text="Patient eupnéique")
    examen.click()
    page.keyboard.press("Control+Home")
    page.keyboard.press("End")
    page.keyboard.press("Shift+Control+ArrowLeft")
    page.keyboard.press("Shift+Control+ArrowLeft")  # « apyrétique. » : le mot et le point
    page.click('[data-action="gras"]')
    expect(page.locator('[data-action="gras"]')).to_have_attribute("aria-pressed", "true")
    page.keyboard.press("Control+i")
    page.locator("#feuille p.bloc", has_text="Motif").click()  # quitter le paragraphe
    expect(page.locator("#feuille p.bloc", has_text="Patient eupnéique").locator("b i, i b")).to_have_count(1)

    # Cellule de tableau
    page.locator("#feuille td p", has_text="Troponine").click()
    page.keyboard.press("End")
    page.keyboard.type(" Tn")

    page.click("#enregistrer")
    expect(page.locator("#etat")).to_have_text("Enregistré")
    page.screenshot(path=f"{SP}/retouche.png", full_page=True)
    with page.expect_download() as telechargement:
        page.click("#telecharger")
    from docx import Document
    doc = Document(telechargement.value.path())
    textes = [p.text for p in doc.paragraphs]
    assert any("(retouche)Nouveau paragraphe saisi" in t for t in textes), textes
    assert any("dyspnée d'effort" in t for t in textes), textes
    assert doc.tables[0].cell(0, 0).text == "Troponine Tn"
    examen_docx = next(p for p in doc.paragraphs if p.text.startswith("Patient eupnéique"))
    styles = [(r.text, bool(r.bold), bool(r.italic)) for r in examen_docx.runs if r.text]
    assert styles[0] == ("Patient eupnéique, ", False, False) and styles[1][1:] == (True, True), styles
    # En-têtes et pied de page du document d'origine conservés
    section = doc.sections[0]
    assert section.different_first_page_header_footer
    assert "HÔPITAL FICTIF" in section.first_page_header.paragraphs[0].text
    assert "DUPONT Jean" in section.header.paragraphs[0].text
    assert "Hôpital fictif" in section.footer.paragraphs[0].text
    print("10. retouche : saisie, paragraphes, annulation, fusion, remplacement, gras/italique, tableau, .docx, en-têtes")

    page.click("text=← Retour au résultat")
    expect(page.locator(".pastille", has_text="retouché à la main")).to_be_visible()
    print("11. document marqué « retouché à la main » dans le résultat")

    nav.close()

# 422 (saisie refusée, étape 5) est attendu.
inattendues = [e for e in erreurs if "status of 422" not in e]
print("Erreurs JavaScript inattendues :", inattendues or "aucune")
sys.exit(1 if inattendues else 0)
